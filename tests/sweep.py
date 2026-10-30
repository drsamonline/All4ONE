"""Run every runnable tool against generated fixtures and report failures.

This is the gate that keeps the catalogue honest: a tool that cannot run is
either fixed or removed. Network, Windows-only, destructive and
externally-dependent tools are skipped.
"""

from __future__ import annotations

import json
import signal
import sqlite3
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core import deps, registry  # noqa: E402
from core.runner import execute  # noqa: E402
from core.tool import DIR, FILE, FILES, SAVE, Tool  # noqa: E402

FIX = Path(__file__).resolve().parent / "fixtures"
SKIP_CATEGORIES = {"Network"}
SKIP_TOOLS = {
    # need a real repository, an encrypted file, a token or a live service
    "git-status", "git-log", "git-branches", "git-diff", "file-decrypt",
    "jwt-decode", "disk-speed-test", "process-kill", "service-control",
    "registry-write", "registry-delete", "power-actions", "temp-clean",
    "recycle-bin", "defender-scan", "pdf-decrypt", "pdf-encrypt",
    "hosts-file", "folder-shred", "file-shred", "batch-delete-list",
    "env-variables", "public-ip", "whois-lookup", "dns-lookup", "dns-records", "email-dns-check",
    "ping-host", "traceroute", "port-check", "port-scan", "http-get",
    "http-headers", "http-download", "url-status-check", "tls-certificate",
    "netstat-connections", "local-interfaces", "wake-on-lan", "media-info",
}


class Timeout(Exception):
    pass


def _alarm(_signum, _frame):
    raise Timeout()


def make_fixtures() -> None:
    FIX.mkdir(parents=True, exist_ok=True)
    (FIX / "folder").mkdir(exist_ok=True)
    (FIX / "folder" / "sub").mkdir(exist_ok=True)
    (FIX / "folder" / "one.txt").write_text("alpha beta gamma\nsecond line\n", encoding="utf-8")
    (FIX / "folder" / "two.txt").write_text("alpha beta gamma\nsecond line\n", encoding="utf-8")
    (FIX / "folder" / "sub" / "three.log").write_text("x y z\n", encoding="utf-8")

    (FIX / "sample.txt").write_text(
        "The quick brown fox jumps over the lazy dog.\n"
        "Contact: test@example.com or +91 98765 43210.\n"
        "TODO: check this later.\n", encoding="utf-8")
    (FIX / "sample.md").write_text("# Title\n\n- one\n- two\n\n[link](https://example.com)\n",
                                   encoding="utf-8")
    (FIX / "sample.html").write_text(
        "<html><body><h1>Hi</h1><p>Hello <b>world</b>.</p>"
        "<a href='https://example.com'>link</a></body></html>", encoding="utf-8")
    (FIX / "sample.yaml").write_text("name: test\nitems:\n  - a\n  - b\n", encoding="utf-8")
    (FIX / "sample.ini").write_text("[main]\nkey = value\nother = 2\n", encoding="utf-8")
    (FIX / "sample.py").write_text(
        "import os\n\n\nclass Thing:\n    \"\"\"A thing.\"\"\"\n\n    def go(self):\n"
        "        return os.name\n\n\ndef helper():\n    return 1\n", encoding="utf-8")
    (FIX / "sample.sql").write_text(
        "select a,b from t where a=1 group by a order by b", encoding="utf-8")
    (FIX / "sample.css").write_text("/* c */\n.a { color : red ; }\n", encoding="utf-8")
    (FIX / "sample.js").write_text("// note\nfunction a() { return 1; }\n", encoding="utf-8")
    (FIX / "sample.env").write_text("API_KEY=abc\nEMPTY=\nAPI_KEY=dup\n", encoding="utf-8")
    (FIX / "sample.srt").write_text("1\n00:00:01,000 --> 00:00:02,000\nHello\n", encoding="utf-8")
    (FIX / "requirements.txt").write_text("pypdf>=3\nnotarealpackage123\n", encoding="utf-8")
    (FIX / "list.txt").write_text(f"{FIX / 'sample.txt'}\n{FIX / 'sample.csv'}\n", encoding="utf-8")
    (FIX / "sample.csv").write_text(
        "name,city,amount\nAarav,Mumbai,1200\nDiya,Delhi,800\nKabir,Pune,1500\n", encoding="utf-8")
    (FIX / "other.csv").write_text(
        "name,city,amount\nAarav,Mumbai,1300\nMeera,Chennai,900\n", encoding="utf-8")
    (FIX / "sample.json").write_text(json.dumps(
        [{"id": 1, "name": "A", "tags": ["x"]}, {"id": 2, "name": "B", "tags": []}]),
        encoding="utf-8")
    (FIX / "sample.jsonl").write_text('{"a": 1}\n{"a": 2}\n', encoding="utf-8")
    (FIX / "sample.xml").write_text(
        "<root><item id='1'>one</item><item id='2'>two</item></root>", encoding="utf-8")

    connection = sqlite3.connect(FIX / "sample.db")
    connection.execute("CREATE TABLE IF NOT EXISTS people (id INTEGER, name TEXT)")
    connection.execute("DELETE FROM people")
    connection.executemany("INSERT INTO people VALUES (?, ?)", [(1, "A"), (2, "B")])
    connection.commit()
    connection.close()

    with zipfile.ZipFile(FIX / "sample.zip", "w") as archive:
        archive.writestr("inner.txt", "hello")
        archive.writestr("nested/inner2.txt", "world")

    try:
        from PIL import Image
        image = Image.new("RGB", (64, 48), (30, 60, 120))
        image.save(FIX / "sample.png")
        image.save(FIX / "sample.jpg", quality=90)
        image.save(FIX / "folder" / "pic1.png")
        image.save(FIX / "folder" / "pic2.png")
    except ImportError:
        pass

    try:
        import tarfile
        with tarfile.open(FIX / "sample.tar.gz", "w:gz") as archive:
            archive.add(FIX / "sample.txt", arcname="sample.txt")
    except OSError:
        pass

    try:
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.add_blank_page(width=200, height=200)
        with open(FIX / "sample.pdf", "wb") as handle:
            writer.write(handle)
    except ImportError:
        pass

    try:
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(["name", "amount"])
        sheet.append(["Aarav", 1200])
        book.save(FIX / "sample.xlsx")
    except ImportError:
        pass

    (FIX / "folder" / "data.json").write_text('[{"a": 1}]', encoding="utf-8")
    (FIX / "sample.txt.part001").write_text("alpha ", encoding="utf-8")
    (FIX / "sample.txt.part002").write_text("beta", encoding="utf-8")
    (FIX / "folder" / "doc.pdf").write_bytes((FIX / "sample.pdf").read_bytes()) \
        if (FIX / "sample.pdf").exists() else None

    for name in ("sample.pdf", "sample.csv", "sample.xlsx", "sample.docx", "sample.json",
                 "sample.zip", "sample.png"):
        source = FIX / name
        twin = FIX / name.replace("sample.", "sample2.")
        if source.exists() and not twin.exists():
            twin.write_bytes(source.read_bytes())

    try:
        import docx
        document = docx.Document()
        document.add_paragraph("Hello from a test document.")
        document.add_paragraph("Second paragraph.")
        document.save(FIX / "sample.docx")
    except ImportError:
        pass


def file_fixture(field, tool: Tool) -> str:
    text = (field.filters or "").lower()
    category = tool.category
    if tool.id == "tar-extract":
        return str(FIX / "sample.tar.gz")
    if tool.id == "join-files":
        return str(FIX / "sample.txt.part001")
    if tool.id == "json-extract" and field.key == "file":
        return ""
    if "pdf" in text or category == "PDF":
        return str(FIX / "sample.pdf")
    if "image" in text or "png" in text or "jpg" in text or category == "Images":
        return str(FIX / "sample.png")
    if "excel" in text or "xlsx" in text:
        return str(FIX / "sample.xlsx")
    if "word" in text or "docx" in text:
        return str(FIX / "sample.docx")
    if "csv" in text or "tsv" in text:
        return str(FIX / "sample.csv")
    if "json" in text:
        return str(FIX / "sample.json")
    if "xml" in text:
        return str(FIX / "sample.xml")
    if "yaml" in text:
        return str(FIX / "sample.yaml")
    if "markdown" in text or "*.md" in text:
        return str(FIX / "sample.md")
    if "html" in text:
        return str(FIX / "sample.html")
    if "zip" in text or category == "Archives":
        return str(FIX / "sample.zip")
    if "database" in text or ".db" in text:
        return str(FIX / "sample.db")
    if "python" in text:
        return str(FIX / "sample.py")
    if "config" in text or ".ini" in text:
        return str(FIX / "sample.ini")
    if "subtitle" in text:
        return str(FIX / "sample.srt")
    if field.key in ("second", "file_b", "other", "second_file"):
        return str(FIX / "other.csv") if category == "Data" else str(FIX / "sample.png") \
            if category == "Images" else str(FIX / "sample.pdf") \
            if category == "PDF" else str(FIX / "sample.txt")
    if tool.id in ("tar-extract",):
        return str(FIX / "sample.tar.gz")
    if tool.id in ("xml-format", "xml-to-json"):
        return str(FIX / "sample.xml")
    if tool.id in ("yaml-to-json", "yaml-validate"):
        return str(FIX / "sample.yaml")
    if tool.id == "base64-decode":
        return ""
    return str(FIX / "sample.txt")


def text_fixture(key: str, tool: Tool) -> str:
    key = key.lower()
    if tool.id == "base64-decode":
        return "aGVsbG8gd29ybGQ="
    if tool.id in ("hash-identify",):
        return "d41d8cd98f00b204e9800998ecf8427e"
    if tool.id in ("card-validate",):
        return "4111111111111111"
    if tool.id in ("totp-generate", "totp-verify"):
        return "JBSWY3DPEHPK3PXP" if key == "secret" else "123456"
    if tool.id in ("uuid-generate", "nanoid"):
        return ""
    if key in ("delimiter",):
        return ","
    if key in ("modified", "timestamp"):
        return "2024-01-01 09:30"
    if key in ("start", "end", "from_date", "to_date", "birthdate", "date", "modified",
               "start_date", "end_date"):
        return "2024-01-01"
    if key in ("when", "start_time"):
        return "2024-01-01 09:00"
    if key in ("pages", "order"):
        return "1" if key == "pages" else "2,1"
    if key in ("numbers", "values", "list"):
        return "12, 18, 24"
    if key == "ratio":
        return "16:9"
    if key == "version":
        return "1.4.2"
    if key in ("other",):
        return "1.5.0"
    if key in ("path",):
        return "0.name"
    if key in ("number",):
        return "255"
    if key in ("holidays",):
        return ""
    if key in ("from_zone",):
        return "IST"
    if key in ("to_zone",):
        return "UTC"
    if key == "key" and tool.id == "csv-compare":
        return "name"
    if tool.id == "epoch-convert":
        return "2024-01-01"
    if tool.id == "roman":
        return "2026"
    if key in ("expected", "hash"):
        return "d41d8cd98f00b204e9800998ecf8427e"
    if key in ("expression", "cron"):
        return "0 9 * * 1-5"
    if key == "mac":
        return "AA:BB:CC:DD:EE:FF"
    if key in ("secret", "key"):
        return "JBSWY3DPEHPK3PXP" if key == "secret" else "test-key"
    if key in ("from_unit",):
        return "km"
    if key in ("to_unit",):
        return "mile"
    if key in ("unit",):
        return "m"
    if key in ("table",):
        return "people"
    if key in ("sheet",):
        return "Sheet"
    if key in ("column", "sort_by", "key_column", "keep", "drop", "columns", "pattern_columns"):
        return "name"
    if key in ("rename",):
        return "amount=total"
    if key in ("find", "needle", "pattern", "tags", "text", "value", "replace"):
        return {"find": "alpha", "needle": "alpha", "pattern": "alpha", "tags": "TODO,FIXME",
                "text": "alpha beta gamma", "value": "alpha", "replace": "omega"}.get(
                    key, "alpha")
    if key in ("separator", "prefix", "suffix", "label", "title", "name", "root"):
        return "x"
    if key == "notation":
        return "2d6"
    if key == "sql":
        return "SELECT * FROM people"
    if key == "schema":
        return '{"type": "array"}'
    if key in ("ratio", "aspect"):
        return "1.5"
    if key in ("url", "host", "domain", "address", "cidr", "network"):
        return "example.com"
    return "alpha beta gamma"


def values_for(tool: Tool) -> dict:
    values: dict = {}
    for field in tool.fields:
        if field.kind in (FILE, FILES):
            one = file_fixture(field, tool)
            if field.kind == FILES and one:
                second = one.replace("sample.", "sample2.")
                if Path(second).exists():
                    one = f"{one}|{second}"
            values[field.key] = one
        elif field.kind == DIR:
            if field.key in ("destination", "dest", "target", "output", "backup", "folder_b"):
                target_dir = FIX / f"dest-{tool.id}"
                target_dir.mkdir(parents=True, exist_ok=True)
                values[field.key] = str(target_dir)
            else:
                values[field.key] = str(FIX / "folder")
        elif field.kind == SAVE:
            suffix = ".bin"
            text = (field.filters or "").lower()
            for candidate in (".pdf", ".zip", ".csv", ".json", ".xlsx", ".docx", ".txt",
                              ".png", ".jpg", ".ico", ".reg", ".db", ".md", ".yaml"):
                if candidate in text:
                    suffix = candidate
                    break
            values[field.key] = str(FIX / f"out-{tool.id}{suffix}")
        elif field.kind == "text":
            values[field.key] = text_fixture(field.key, tool)
        elif field.kind == "multiline":
            if tool.id in ("json-extract", "json-format", "json-validate"):
                values[field.key] = '[{"name": "A", "id": 1}]'
            elif tool.id in ("json-schema-lite",):
                values[field.key] = '{"type": "array", "items": {"type": "object"}}'
            else:
                values[field.key] = {
                    "sql": "SELECT * FROM people",
                    "schema": '{"type": "array", "items": {"type": "object"}}',
                }.get(field.key, "alpha beta gamma\nsecond line\n")
        elif field.kind == "password":
            values[field.key] = ("JBSWY3DPEHPK3PXP" if field.key == "secret"
                                 else "Correct-Horse-Battery-9")
        elif field.kind == "color":
            values[field.key] = "#3366cc"
    return values


def main() -> int:
    make_fixtures()
    tools = [t for t in registry.all_tools()
             if t.category not in SKIP_CATEGORIES
             and t.id not in SKIP_TOOLS
             and not t.destructive
             and not t.is_windows_only
             and not deps.missing(t.needs)]
    print(f"running {len(tools)} tools\n")
    failures: list[tuple[str, str]] = []
    signal.signal(signal.SIGALRM, _alarm)

    for tool in tools:
        signal.alarm(25)
        try:
            result = execute(tool, values_for(tool), sink=lambda line: None)
        except Timeout:
            failures.append((tool.id, "timed out after 25s"))
            continue
        finally:
            signal.alarm(0)
        if not result.ok:
            failures.append((tool.id, result.error or "unknown"))

    print(f"passed: {len(tools) - len(failures)} / {len(tools)}")
    if failures:
        print(f"\nfailures ({len(failures)}):")
        for tool_id, error in failures:
            print(f"  {tool_id:<28} {error[:110]}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
